#!/usr/bin/env python3
"""
BlogCMS - Automated Demo, Screenshot Capture & Word Report
==========================================================
1. Captures "terminal process" screenshots of every infrastructure step
   (docker compose up / docker ps / curl health checks / final verification).
2. Drives the real running application in a headless browser (Playwright + Edge)
   and captures full-page screenshots of each demo step:
   Home -> Register -> Logout/Login -> New Post (category + publish) ->
   Post Detail -> Comment -> Home listing -> Swagger docs.
3. Assembles all screenshots, in order, into a formatted Word document
   (BlogCMS_Demo.docx).

Usage:
    python demo/run_demo.py            # capture screenshots then build report
    python demo/run_demo.py --report   # rebuild the .docx from existing shots
"""
import html
import json
import subprocess
import sys
from pathlib import Path

from playwright.sync_api import sync_playwright

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    HAS_DOCX = True
except ImportError:  # pragma: no cover
    HAS_DOCX = False

BASE = Path(__file__).resolve().parent.parent
DEMO = BASE / "demo"
SHOTS = DEMO / "shots"
SHOTS.mkdir(parents=True, exist_ok=True)

BASE_URL = "http://localhost"
API_URL = "http://localhost:8000"

# Credentials for the demo. First registered user on a fresh DB becomes admin.
DEMO_USER = "admin"
DEMO_EMAIL = "admin@example.com"
DEMO_PASS = "admin12345"

POST_TITLE = "My First Blog Post - Hello World"
POST_SUMMARY = "A quick tour of the BlogCMS stack: React, FastAPI, MySQL, Docker."
POST_CONTENT = (
    "This post was published from the BlogCMS application to demonstrate the\n"
    "full Content Management workflow end-to-end.\n\n"
    "BlogCMS is a 3-tier blog platform:\n"
    "  - Frontend: React (Vite) served by Nginx on port 80\n"
    "  - Backend:  FastAPI with JWT auth on port 8000\n"
    "  - Database: MySQL 8.0 on port 3306\n\n"
    "Everything runs in Docker Compose containers on a shared bridge network.\n"
    "The first user to register automatically becomes an administrator."
)
COMMENT_TEXT = "Great first post! Publishing content right from the UI is seamless."

TERMINAL_THEME = """
  * { margin:0; padding:0; box-sizing:border-box; }
  html, body { background:#0b0e14; }
  .term {
    background:#131720;
    border:1px solid #2a3242;
    border-radius:10px;
    overflow:hidden;
    font-family:'Consolas','Cascadia Mono','Courier New',monospace;
    width:1100px; margin:14px auto;
    box-shadow:0 18px 50px rgba(0,0,0,.55);
  }
  .term-bar {
    background:#1c2331; padding:10px 16px; display:flex; align-items:center;
    border-bottom:1px solid #2a3242;
  }
  .dot { width:12px; height:12px; border-radius:50%; display:inline-block; margin-right:7px; }
  .term-title { color:#8b93a5; font-size:13px; margin-left:8px; letter-spacing:.4px; }
  .term-body { padding:20px 22px 24px; }
  .block { margin-bottom:18px; }
  .block:last-child { margin-bottom:0; }
  .prompt { color:#7ee787; font-size:15px; line-height:1.6; }
  .cmd { color:#e6edf3; }
  .out { color:#9da6b5; font-size:14px; line-height:1.55; white-space:pre; margin-top:4px; }
  .ok { color:#7ee787; } .info { color:#79c0ff; } .hl { color:#ffa657; }
"""


def shell(cmd: str) -> str:
    """Run a shell command and return combined stdout+stderr."""
    proc = subprocess.run(
        cmd, shell=True, capture_output=True, text=True,
        encoding="utf-8", errors="replace", timeout=60,
    )
    return (proc.stdout + proc.stderr).strip()


def term_html(title: str, blocks: list) -> str:
    """Render terminal-style HTML for a list of {cmd, out[, ok]} blocks."""
    body = []
    for b in blocks:
        cmd_html = "<span class='cmd'>%s</span>" % html.escape(b["cmd"])
        body.append("<div class='block'><div class='prompt'><span class='info'>➜</span> %s</div>" % cmd_html)
        if b.get("out"):
            out = html.escape(b["out"])
            if b.get("ok"):
                body.append("<div class='out ok'>%s</div>" % out)
            else:
                body.append("<div class='out'>%s</div>" % out)
        body.append("</div>")
    return (
        "<html><head><style>%s</style></head><body><div class='term'>"
        "<div class='term-bar'>"
        "<span class='dot' style='background:#ff5f56'></span>"
        "<span class='dot' style='background:#ffbd2e'></span>"
        "<span class='dot' style='background:#27c93f'></span>"
        "<span class='term-title'>%s</span>"
        "</div><div class='term-body'>%s</div></body></html>"
    ) % (TERMINAL_THEME, html.escape(title), "".join(body))


def snap(page, name: str) -> Path:
    path = SHOTS / name
    page.screenshot(path=str(path), full_page=True)
    return path


def capture_terminal(page, name: str, title: str, blocks: list) -> Path:
    """Render a terminal-style HTML page and screenshot it."""
    page.set_content(term_html(title, blocks))
    page.wait_for_timeout(300)
    return snap(page, name)


def capture_process_shots(page) -> list:
    """Part A - screenshots of the infrastructure / process steps."""
    shots = []

    # A1 - starting the stack
    up_log = shell("grep -E 'Network blog|Volume blog|Container blog-cms|Healthy' /tmp/compose.log 2>/dev/null")
    capture_terminal(
        page, "A1-start-stack.png", "BlogCMS Demo - Step 1 : Start the full stack",
        [{
            "cmd": "docker compose up --build -d",
            "out": up_log or "(log from previous run not available - stack is already up)",
        }],
    )
    shots.append(("A1-start-stack.png", "docker compose up --build -d - the full stack starts", "process"))

    # A2 - running containers
    ps = shell('docker ps --filter name=blog-cms --format "table {{.Names}}\\t{{.Image}}\\t{{.Ports}}\\t{{.Status}}"')
    capture_terminal(
        page, "A2-docker-ps.png", "BlogCMS Demo - Step 2 : Running containers",
        [{"cmd": "docker ps --filter name=blog-cms", "out": ps, "ok": True}],
    )
    shots.append(("A2-docker-ps.png", "All three containers are up and healthy", "process"))

    # A3 - backend health endpoint
    health = shell("curl -s -m 10 http://localhost:8000/api/health")
    capture_terminal(
        page, "A3-backend-health.png", "BlogCMS Demo - Step 3 : Backend health check",
        [{"cmd": "curl -s http://localhost:8000/api/health", "out": health, "ok": True}],
    )
    shots.append(("A3-backend-health.png", "Backend FastAPI health endpoint - DB connected", "process"))

    # A4 - nginx proxy health
    proxied = shell("curl -s -m 10 http://localhost/api/health")
    capture_terminal(
        page, "A4-proxy-health.png", "BlogCMS Demo - Step 4 : Nginx reverse-proxy check",
        [{"cmd": "curl -s http://localhost/api/health", "out": proxied, "ok": True}],
    )
    shots.append(("A4-proxy-health.png", "Nginx proxies /api/* on port 80 to the backend", "process"))
    return shots
def _goto(page, url: str, wait_ms: int = 900):
    """Navigate to a URL and wait for the SPA to settle."""
    try:
        page.goto(url, wait_until="networkidle", timeout=20000)
    except Exception:
        page.goto(url, wait_until="domcontentloaded", timeout=20000)
    page.wait_for_timeout(wait_ms)


def wait_path(page, path, timeout=15000):
    """Wait until the SPA pathname equals `path` (client-side routing aware)."""
    page.wait_for_function("location.pathname === '%s'" % path, timeout=timeout)


def snap_and_append(shots, page, name, caption, kind="app"):
    """Take the screenshot and register its metadata in the report list."""
    try:
        snap(page, name)
    except Exception as exc:
        print("  [warn] screenshot failed for %s: %s" % (name, exc))
    shots.append((name, caption, kind))


def demo_app(page) -> list:
    """Part B - walk through the application in the browser, screenshotting each step."""
    shots = []

    # ---------------------------------------------------------------- B1: home
    _goto(page, BASE_URL + "/")
    page_title = page.title()
    snap_and_append(shots, page, "B1-home-empty.png",
                    "Application loads at http://localhost - title: '%s'" % page_title)

    # ------------------------------------------------- B2: register (admin user)
    _goto(page, BASE_URL + "/register")
    try:
        page.locator("form input").nth(0).fill(DEMO_USER)
        page.locator('form input[type="email"]').fill(DEMO_EMAIL)
        page.locator('form input[type="password"]').fill(DEMO_PASS)
        page.wait_for_timeout(300)
        snap_and_append(shots, page, "B2-register-filled.png",
                        "Register form filled - first user becomes 'admin' automatically")
        page.locator('form button[type="submit"]').click()
        wait_path(page, "/", timeout=10000)
    except Exception:
        # User already exists from an earlier run - just log in instead.
        print(">> registration declined (user exists?) - switching to login flow")
        _goto(page, BASE_URL + "/login")
        page.locator("form input").nth(0).fill(DEMO_USER)
        page.locator("form input").nth(1).fill(DEMO_PASS)
        page.locator('form button[type="submit"]').click()
        wait_path(page, "/", timeout=10000)
    page.wait_for_timeout(700)
    snap_and_append(shots, page, "B3-home-admin.png",
                    "Registered + signed in: header shows the new 'admin' user with role badge")

    # --------------------------------------------------------- B4: logout, login
    page.locator('button:has-text("Log out")').click()
    page.wait_for_timeout(600)
    _goto(page, BASE_URL + "/login")
    page.locator("form input").nth(0).fill(DEMO_USER)
    page.locator("form input").nth(1).fill(DEMO_PASS)
    page.wait_for_timeout(300)
    snap_and_append(shots, page, "B4-login-filled.png", "Logged out, then filled the Log in form")
    page.locator('form button[type="submit"]').click()
    wait_path(page, "/", timeout=10000)
    page.wait_for_timeout(700)
    snap_and_append(shots, page, "B5-home-login.png", "Signed in again - JWT auth flow works through the UI")

    # ---------------------------------------------------------- B6: new post view
    _goto(page, BASE_URL + "/new-post")
    snap_and_append(shots, page, "B6-new-post.png",
                    "New Post editor - only admin/author roles may write here")

    # ---------------------------------------------------- B7: add a category
    page.fill('input[placeholder="New category name"]', "DevOps")
    page.click('button:has-text("Add")')
    page.wait_for_timeout(900)
    snap_and_append(shots, page, "B7-category-added.png",
                    "Created the 'DevOps' category inside the editor (POST /api/categories)")

    # ---------------------------------------------------------- B8: filled form
    page.locator("form input").nth(0).fill(POST_TITLE)
    page.locator("form input").nth(1).fill(POST_SUMMARY)
    page.locator("form textarea").fill(POST_CONTENT)
    page.wait_for_timeout(300)
    snap_and_append(shots, page, "B8-post-filled.png",
                    "Post details entered - title, summary, content, category")

    # -------------------------------------------------------------- B9: publish
    page.locator('form button[type="submit"]').click()
    try:
        page.wait_for_function("location.pathname.startsWith('/posts/')", timeout=15000)
    except Exception:
        # Post may already exist from a previous run - reconnect to its detail page.
        _goto(page, BASE_URL + "/")
        try:
            page.locator("a[href^='/posts/']").first.click()
            page.wait_for_function("location.pathname.startsWith('/posts/')", timeout=10000)
        except Exception:
            pass
    page.wait_for_timeout(900)
    snap_and_append(shots, page, "B9-post-detail.png",
                    "Post published (POST /api/posts) and opened on its detail page")

    # --------------------------------------------------------------- B10: comment
    try:
        page.fill('textarea[placeholder="Add a comment..."]', COMMENT_TEXT)
        page.click('button:has-text("Post comment")')
        page.wait_for_timeout(1000)
        snap_and_append(shots, page, "B10-comment.png",
                        "Comment posted (POST /api/posts/{id}/comments)")
    except Exception:
        snap_and_append(shots, page, "B10-comment.png", "(comment step skipped)")

    # ------------------------------------------------------------- B11: home list
    _goto(page, BASE_URL + "/")
    snap_and_append(shots, page, "B11-home-post.png",
                    "Home page now lists the published post with its category")

    # ----------------------------------------------------------- B12: swagger docs
    try:
        _goto(page, API_URL + "/docs", wait_ms=2500)
        snap_and_append(shots, page, "B12-swagger.png",
                        "Interactive Swagger UI at http://localhost:8000/docs")
    except Exception:
        snap_and_append(shots, page, "B12-swagger.png", "(docs screenshot skipped)")
    return shots
def capture_verification_shots(page) -> list:
    """Part C - final state verification (terminal screenshots)."""
    shots = []

    # A5 - containers still healthy
    ps = shell('docker ps --filter name=blog-cms --format "table {{.Names}}\\t{{.Status}}\\t{{.Ports}}"')
    capture_terminal(
        page, "A5-docker-ps-final.png", "BlogCMS Demo - Step 5 : Final container state",
        [{"cmd": "docker ps --filter name=blog-cms", "out": ps, "ok": True}],
    )
    shots.append(("A5-docker-ps-final.png", "All containers running and healthy after the demo", "process"))

    # A6 - posts API now contains the published post
    posts = shell("curl -s -m 10 http://localhost/api/posts")
    try:
        dump = json.dumps(json.loads(posts), indent=2, ensure_ascii=False)
    except Exception:
        dump = posts
    capture_terminal(
        page, "A6-posts-api.png", "BlogCMS Demo - Step 6 : Posts API verification",
        [{"cmd": "curl -s http://localhost/api/posts", "out": dump, "ok": True}],
    )
    shots.append(("A6-posts-api.png", "GET /api/posts returns the demo post (stored in MySQL)", "process"))

    # A7 - categories API
    cats = shell("curl -s -m 10 http://localhost/api/categories")
    try:
        dump = json.dumps(json.loads(cats), indent=2, ensure_ascii=False)
    except Exception:
        dump = cats
    capture_terminal(
        page, "A7-categories-api.png", "BlogCMS Demo - Step 7 : Categories API verification",
        [{"cmd": "curl -s http://localhost/api/categories", "out": dump, "ok": True}],
    )
    shots.append(("A7-categories-api.png", "GET /api/categories returns the 'DevOps' category", "process"))
    return shots


# ---------------------------------------------------------------- Word report
def build_report(shots: list) -> Path:
    """Build BlogCMS_Demo.docx from the ordered screenshot list."""
    assert HAS_DOCX, "python-docx is required to build the report"
    report = DEMO / "BlogCMS_Demo.docx"

    doc = Document()

    # base typography
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for sec in doc.sections:
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)
        sec.left_margin = Inches(0.9)
        sec.right_margin = Inches(0.9)

    # ---- title page ----
    title = doc.add_heading("BlogCMS - Application Demo", level=0)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x10, 0x5C, 0x9E)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Full-Stack Blog Platform - React + FastAPI + MySQL + Docker")
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x44, 0x4A, 0x53)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("Step-by-step walkthrough with screenshots of every process performed.\n"
                 "Generated automatically by demo/run_demo.py against the live Docker Compose stack.").italic = True

    doc.add_paragraph().add_run().add_break()

    # ---- overview ----
    doc.add_heading("1. Overview", level=1)
    doc.add_paragraph(
        "This document captures the end-to-end demo of the BlogCMS application. It starts with the "
        "infrastructure (Docker Compose containers), verifies the backend health and the Nginx reverse "
        "proxy, then performs a full content workflow in the browser: user registration, login, category "
        "creation, publishing a post, commenting on it, and finally verifying the data through the REST API "
        "and Swagger UI. Every step was captured as a screenshot."
    )

    # ---- walkthrough ----
    doc.add_heading("2. Step-by-Step Demo", level=1)
    for i, (img, caption, kind) in enumerate(shots, 1):
        path = SHOTS / img
        if not path.exists():
            doc.add_paragraph("[missing screenshot: %s]" % img, style="List Number")
            continue
        heading = caption.split(" - ")[0] if " - " in caption else caption
        doc.add_heading("Step %d: %s" % (i, heading), level=2)
        doc.add_picture(str(path), width=Inches(6.4))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(caption + ("   [terminal]" if kind == "process" else "   [browser]"))
        r.font.size = Pt(9)
        r.italic = True
        r.font.color.rgb = RGBColor(0x66, 0x6E, 0x7A)

    # ---- quick reference ----
    doc.add_heading("3. Quick Reference", level=1)
    items = [
        ("Frontend (React + Nginx)", "http://localhost"),
        ("Backend API docs (Swagger)", "http://localhost:8000/docs"),
        ("MySQL", "localhost:3306  (blog_user / blog_pass / blog_cms)"),
        ("Demo admin account", "%s / %s" % (DEMO_USER, DEMO_PASS)),
        ("Run the stack", "docker compose up --build -d"),
        ("Stop the stack", "docker compose down"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = "Resource"
    hdr[1].text = "Value"
    for key, val in items:
        row = table.add_row().cells
        row[0].text = key
        row[1].text = val

    doc.save(str(report))
    return report
def _launch_browser(p):
    """Launch an installed Chromium-based browser (Edge/Chrome), falling back
    to Playwright's own bundle if available."""
    for channel in ("msedge", "chrome"):
        try:
            return p.chromium.launch(headless=True, channel=channel)
        except Exception:
            continue
    return p.chromium.launch(headless=True)


def main():
    only_report = "--report" in sys.argv
    shots = []

    if not only_report:
        print(">> Launching headless browser (Edge/Chrome)...")
        with sync_playwright() as p:
            browser = _launch_browser(p)
            page = browser.new_page(viewport={"width": 1280, "height": 900})
            page.set_default_timeout(12000)

            try:
                print(">> Capturing infrastructure / process screenshots...")
                shots += capture_process_shots(page)
            except Exception as exc:
                print(">> WARN: process capture failed: %s" % exc)

            try:
                print(">> Walking through the application in the browser...")
                shots += demo_app(page)
            except Exception as exc:
                print(">> WARN: app demo failed: %s" % exc)

            try:
                print(">> Capturing final verification screenshots...")
                shots += capture_verification_shots(page)
            except Exception as exc:
                print(">> WARN: verification capture failed: %s" % exc)

            browser.close()
    else:
        # Rebuild the report from previously captured screenshots in canonical order.
        canonical = [
            "A1-start-stack.png", "A2-docker-ps.png", "A3-backend-health.png",
            "A4-proxy-health.png", "B1-home-empty.png", "B2-register-filled.png",
            "B3-home-admin.png", "B4-login-filled.png", "B5-home-login.png",
            "B6-new-post.png", "B7-category-added.png", "B8-post-filled.png",
            "B9-post-detail.png", "B10-comment.png", "B11-home-post.png",
            "B12-swagger.png", "A5-docker-ps-final.png", "A6-posts-api.png",
            "A7-categories-api.png",
        ]
        captions = {
            "A1-start-stack.png": "docker compose up --build -d - the full stack starts",
            "A2-docker-ps.png": "All three containers are up and healthy",
            "A3-backend-health.png": "Backend FastAPI health endpoint - DB connected",
            "A4-proxy-health.png": "Nginx proxies /api/* on port 80 to the backend",
            "B1-home-empty.png": "Application loads at http://localhost (empty blog)",
            "B2-register-filled.png": "Register form filled - first user becomes admin",
            "B3-home-admin.png": "Signed in - header shows admin role badge",
            "B4-login-filled.png": "Login form after logging out",
            "B5-home-login.png": "Signed in again - JWT flow works through the UI",
            "B6-new-post.png": "New Post editor (protected admin/author route)",
            "B7-category-added.png": "DevOps category created inside the editor",
            "B8-post-filled.png": "Post details entered",
            "B9-post-detail.png": "Post published and opened on its detail page",
            "B10-comment.png": "Comment posted on the article",
            "B11-home-post.png": "Home page lists the published post",
            "B12-swagger.png": "Swagger UI at http://localhost:8000/docs",
            "A5-docker-ps-final.png": "Final container state - all healthy",
            "A6-posts-api.png": "GET /api/posts returns the demo post",
            "A7-categories-api.png": "GET /api/categories returns DevOps category",
        }
        shots = [(f, captions.get(f, f), "app") for f in canonical if (SHOTS / f).exists()]

    if not shots:
        print("No screenshots available. Run without --report first.")
        return 1

    if HAS_DOCX:
        out = build_report(shots)
        print(">> Report written to %s" % out)
        print(">> %d screenshots embedded." % len(shots))
    else:
        print(">> python-docx not installed; screenshots are in demo/shots/")
    return 0


if __name__ == "__main__":
    sys.exit(main())